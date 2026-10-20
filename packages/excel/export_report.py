"""Production progress and activity report exporter to styled Excel."""

from __future__ import annotations
from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def generate_production_excel_report(db_path: str) -> BytesIO:
    """Build the production report from the engine database behind ``db_path``.

    The report used to express its lookups as one joined statement using
    ``GROUP BY``, ``coalesce``, ``||`` and ``case when``. The engine supports none
    of those, and its join merges rows by column name, so every lookup is now a
    flat SELECT resolved in Python.
    """
    from packages.persistence.database import engine_path
    from packages.sqlengine import connect

    database = connect(engine_path(f"sqlite:///{Path(db_path).as_posix()}"))

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "TỔNG QUAN TIẾN ĐỘ"

    # Fetch stats
    total_prods = int(database.scalar("SELECT COUNT(*) FROM products") or 0)
    total_tracking = int(database.scalar("SELECT COUNT(*) FROM order_tracking_items") or 0)
    status_counts: dict[str, int] = {}
    for row in database.query("SELECT status FROM order_tracking_items"):
        key = str(row["status"])
        status_counts[key] = status_counts.get(key, 0) + 1
    total_reports = int(database.scalar("SELECT COUNT(*) FROM process_report_events") or 0)
    total_workflows = int(database.scalar("SELECT COUNT(*) FROM tracking_workflow_events") or 0)

    # Styles
    title_font = Font(name="Segoe UI", size=16, bold=True, color="1B365D")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    sub_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    bold_font = Font(name="Segoe UI", size=11, bold=True)
    normal_font = Font(name="Segoe UI", size=11)
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    # --- SHEET 1: SUMMARY ---
    ws_summary["A1"] = "BÁO CÁO TIẾN ĐỘ SẢN XUẤT & TRUY VẾT QR — HMS FACTORY"
    ws_summary["A1"].font = title_font
    ws_summary["A2"] = f"Thời gian xuất báo cáo: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
    ws_summary["A2"].font = Font(name="Segoe UI", size=10, italic=True, color="6B7280")

    ws_summary["A4"] = "CHỈ SỐ TỔNG HỢP"
    ws_summary["A4"].font = Font(name="Segoe UI", size=12, bold=True, color="1E3A8A")

    kpis = [
        ("Tổng số sản phẩm quản lý", total_prods, "Chi tiết"),
        ("Tổng số mã QR theo dõi (Tracking)", total_tracking, "Mã QR"),
        ("Đang tiến hành gia công (IN_PRODUCTION)", status_counts.get("IN_PRODUCTION", 0) + status_counts.get("IN_PROCESS", 0), "Chi tiết"),
        ("Đã kiểm tra chất lượng (QC_CHECKED)", status_counts.get("QC_CHECKED", 0), "Chi tiết"),
        ("Đã hoàn tất đóng gói (PACKED)", status_counts.get("PACKED", 0), "Chi tiết"),
        ("Đã xuất xưởng giao hàng (DELIVERED)", status_counts.get("DELIVERED", 0), "Chi tiết"),
        ("Tổng lượt báo cáo công đoạn gia công", total_reports, "Lượt"),
        ("Tổng lượt sự kiện quy trình QC & xuất hàng", total_workflows, "Lượt"),
    ]

    ws_summary.append([])
    ws_summary.append(["Chỉ số hoạt động", "Số lượng", "Đơn vị tính"])
    for cell in ws_summary[6]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, (label, val, unit) in enumerate(kpis, start=7):
        ws_summary.append([label, val, unit])
        row_cells = ws_summary[idx]
        row_cells[0].font = normal_font
        row_cells[1].font = bold_font
        row_cells[1].alignment = Alignment(horizontal="right")
        row_cells[2].font = normal_font
        row_cells[2].alignment = Alignment(horizontal="center")
        for c in row_cells:
            c.border = thin_border

    # --- SHEET 2: TRACKING ITEMS ---
    ws_items = wb.create_sheet(title="DANH SÁCH CHI TIẾT")
    headers_items = [
        "STT", "Mã Tracking", "Mã Sản Phẩm", "Tên Chi Tiết", "Khách Hàng",
        "Số Lượng", "ĐVT", "Mác Vật Liệu", "Quy Cách / Kích Thước",
        "Xử Lý Bề Mặt", "Hạn Giao", "Trạng Thái", "Mã QR Public"
    ]
    ws_items.append(headers_items)
    for cell in ws_items[1]:
        cell.font = header_font
        cell.fill = sub_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    machining_types = {
        row["internal_id"]: row["display_name"]
        for row in database.query("SELECT internal_id, display_name FROM machining_types")
    }
    tracking_items = {
        row["internal_id"]: row
        for row in database.query("SELECT internal_id, tracking_code, product_id FROM order_tracking_items")
    }
    products = {
        row["internal_id"]: row
        for row in database.query(
            "SELECT internal_id, part_name, product_code, company, material, size,"
            " surface_treatment FROM products"
        )
    }

    items_rows = []
    for item in database.query(
        "SELECT internal_id, tracking_code, quantity, unit, delivery_date, status,"
        " qr_public_id, product_id FROM order_tracking_items"
    ):
        product = products.get(item["product_id"])
        if product is None:
            continue
        items_rows.append({
            "tracking_code": item["tracking_code"],
            "product_code": product.get("product_code"),
            "part_name": product["part_name"],
            "company": product.get("company"),
            "quantity": item["quantity"],
            "unit": item["unit"],
            "material": product.get("material"),
            "size": product.get("size"),
            "surface_treatment": product.get("surface_treatment"),
            "delivery_date": item["delivery_date"],
            "status": item["status"],
            "qr_public_id": item["qr_public_id"],
        })
    # ``ORDER BY rowid`` is a no-op on the engine, so newest-first is restored here.
    items_rows.reverse()

    for idx, r in enumerate(items_rows, start=1):
        raw_q = r["quantity"]
        try:
            fq = float(raw_q)
            q = int(fq) if fq.is_integer() else fq
        except Exception:
            q = raw_q
        row_data = [
            idx, r["tracking_code"], r["product_code"], r["part_name"], r["company"],
            q, r["unit"], r["material"] or "", r["size"] or "",
            r["surface_treatment"] or "", str(r["delivery_date"] or ""),
            r["status"], r["qr_public_id"] or ""
        ]
        ws_items.append(row_data)
        for c in ws_items[idx + 1]:
            c.font = normal_font
            c.border = thin_border

    # --- SHEET 3: ACTIVITY LOG ---
    ws_log = wb.create_sheet(title="NHẬT KÝ SẢN XUẤT")
    headers_log = [
        "Thời Gian", "Nhân Sự", "Chi Tiết Sản Phẩm", "Mã Tracking",
        "Công Đoạn / Thao Tác", "Phân Loại", "Số Lượng", "Nội Dung Báo Cáo"
    ]
    ws_log.append(headers_log)
    for cell in ws_log[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    reports = []
    for r in database.query(
        "SELECT server_timestamp, actor_display_name_snapshot, machining_type_id,"
        " attempt_number, tracking_item_id, quantity, notes FROM process_report_events"
    ):
        tracking = tracking_items.get(r["tracking_item_id"]) or {}
        product = products.get(tracking.get("product_id")) or {}
        machining = machining_types.get(r["machining_type_id"]) or "Gia công"
        attempt = r["attempt_number"]
        action = f"{machining} (Lần {attempt})" if attempt is not None else f"{machining} (Đã xong)"
        reports.append({
            "server_timestamp": r["server_timestamp"],
            "actor": r["actor_display_name_snapshot"],
            "part_name": product.get("part_name"),
            "tracking_code": tracking.get("tracking_code"),
            "action": action,
            "category": "GIA CÔNG",
            "quantity": r["quantity"],
            "notes": r["notes"],
        })

    workflows = []
    for w in database.query(
        "SELECT server_timestamp, actor_display_name_snapshot, event_type,"
        " tracking_item_id, quantity, notes FROM tracking_workflow_events"
    ):
        tracking = tracking_items.get(w["tracking_item_id"]) or {}
        product = products.get(tracking.get("product_id")) or {}
        workflows.append({
            "server_timestamp": w["server_timestamp"],
            "actor": w["actor_display_name_snapshot"],
            "part_name": product.get("part_name"),
            "tracking_code": tracking.get("tracking_code"),
            "action": w["event_type"],
            "category": "QUY TRÌNH QC/GIAO",
            "quantity": w["quantity"],
            "notes": w["notes"],
        })

    all_logs = sorted(reports + workflows, key=lambda x: str(x["server_timestamp"]), reverse=True)

    for idx, l in enumerate(all_logs, start=1):
        ts = str(l["server_timestamp"])[:19].replace("T", " ")
        raw_q = l["quantity"]
        q = ""
        if raw_q is not None and str(raw_q).strip():
            try:
                fq = float(raw_q)
                q = int(fq) if fq.is_integer() else fq
            except Exception:
                q = raw_q

        ws_log.append([
            ts, l["actor"], l["part_name"] or "", l["tracking_code"] or "",
            l["action"], l["category"], q, l["notes"] or ""
        ])
        for c in ws_log[idx + 1]:
            c.font = normal_font
            c.border = thin_border

    # Auto-adjust column widths for all sheets
    for ws in [ws_summary, ws_items, ws_log]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 45)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
