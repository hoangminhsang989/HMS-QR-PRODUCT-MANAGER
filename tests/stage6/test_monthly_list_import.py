"""Regression gate for the dated monthly-list import used by real customer workbooks.

Builds a minimal ``LIST HÀNG THÁNG`` workbook with the customer's exact shape
(day sheets, header row 6, row-anchored pictures in column B, merged Size
header) and proves the import is correct and idempotent: products, images,
tracking items, QR codes and repaired product codes all land once and only
once.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as OpenpyxlImage

from config.paths import require_test_root
from packages.excel.monthly_list_import import import_monthly_lists


PNG_1X1 = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01"
    b"\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
)

HEADERS = [
    "Company",
    "Picture",
    "Part Name",
    "Quantity",
    "Unit",
    None,
    None,
    "Material",
    "NGƯỜI NHẬN",
    "Surface treatment",
    None,
    None,
    "Size",
    None,
    None,
    "Ghi chú",
    "Lựa Vendor giao hàng",
]

DAY_03_ROWS = [
    ["CIT", None, "Jig dap nhiet", 2, "SET", 485000, 970000, "AL6061", "MR HIỆP", None, None, None, 85, 55, 30, "Giao tuan sau", None],
    ["CIT", None, "Chi tiet may lanh", 2, "EA", 350000, 700000, "AL6061", "MR HIỆP", None, None, None, 99, 23.25, 12, None, None],
    ["CIT", None, "Ban nang thuy luc", 5, "EA", 100000, 500000, "SUS 304", "MR HIỆP", None, None, None, "Ø", 18, 11.5, None, None],
]

DAY_04_ROWS = [
    ["CIT", None, "Khi nen khi", 3, "EA", 250000, 750000, "DUROSTONE", "MR HIỆP", None, None, None, 40, 16, 3, None, None],
    ["CIT", None, "De nhua dau", 1, "EA", 150000, 150000, "SUS 304", "MR HIỆP", None, None, None, 8, 45, 20, None, None],
]


def _build_workbook(path: Path, image: Path) -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "03-06"
    second = workbook.create_sheet("04-06")
    for sheet, rows in ((first, DAY_03_ROWS), (second, DAY_04_ROWS)):
        sheet["C1"] = "CONG TY TNHH SHL TECH"
        sheet.merge_cells("C1:O1")
        sheet["P1"] = "2026-06-03"
        for column, header in enumerate(HEADERS, start=1):
            if header is not None:
                sheet.cell(row=6, column=column, value=header)
        sheet.merge_cells("M6:O6")
        for row_index, row in enumerate(rows, start=7):
            for column, value in enumerate(row, start=1):
                if value is not None:
                    sheet.cell(row=row_index, column=column, value=value)
    first.add_image(OpenpyxlImage(image), "B7")
    workbook.save(path)
    workbook.close()


def _query(database: Path, sql: str, params: tuple = ()) -> list[tuple]:
    from packages.excel.monthly_list_import import create_database_runtime_for_shared_sqlite

    runtime = create_database_runtime_for_shared_sqlite(database)
    try:
        rows = runtime.database.query(sql, list(params))
    finally:
        runtime.dispose()
    if not rows:
        return []
    return [tuple(row.values()) for row in rows]


def test_monthly_list_import_is_correct_and_idempotent(tmp_path: Path):
    assert require_test_root() in tmp_path.resolve().parents
    source = tmp_path / "LIST HÀNG THÁNG -06-26.xlsx"
    image = tmp_path / "part.png"
    image.write_bytes(PNG_1X1)
    _build_workbook(source, image)
    database = tmp_path / "monthly.sqlite"
    storage = tmp_path / "storage"

    first = import_monthly_lists([source], database_path=database, storage_root=storage)
    assert first["products_created"] == 5
    assert first["products_existing"] == 0
    assert first["images_imported"] == 1
    assert first["images_skipped"] == 0
    assert first["rows_skipped"] == 0
    assert first["tracking_created"] == 5
    assert first["product_codes_repaired"] == 5

    products = _query(
        database,
        "SELECT product_code, company, part_name, quantity, unit, material, size, delivery_schedule "
        "FROM products ORDER BY delivery_schedule, product_code",
    )
    assert len(products) == 5
    assert products[0] == (
        "C-BNTL-260603-R009", "CIT", "Ban nang thuy luc", 5, "EA", "SUS 304", "Ø x 18 x 11.5", date(2026, 6, 3),
    )
    assert products[1][0] == "C-CTML-260603-R008"
    assert products[2] == (
        "C-JDN-260603-R007", "CIT", "Jig dap nhiet", 2, "SET", "AL6061", "85 x 55 x 30", date(2026, 6, 3),
    )
    assert products[3][0] == "C-DND-260604-R008"
    assert products[4][0] == "C-KNK-260604-R007"
    assert {row[7] for row in products} == {date(2026, 6, 3), date(2026, 6, 4)}

    manifest = _query(database, "SELECT source_file, source_sheet, source_row, image_sha256 FROM excel_monthly_import_rows ORDER BY source_row")
    assert len(manifest) == 5
    assert manifest[0][0] == source.name and manifest[0][1] == "03-06" and manifest[0][2] == 7
    assert manifest[0][3] is not None

    tracking = _query(database, "SELECT tracking_code, quantity, unit, delivery_date, status, qr_public_id, qr_status FROM order_tracking_items ORDER BY tracking_code")
    assert len(tracking) == 5
    assert all(row[0].startswith("ITEM-2026-") for row in tracking)
    assert all(row[3] in {date(2026, 6, 3), date(2026, 6, 4)} for row in tracking)
    assert all(row[4] == "NEW" and row[5].startswith("qr_") and row[6] == "ACTIVE" for row in tracking)
    assert _query(database, "SELECT name FROM customers WHERE name = 'CIT'")
    po_numbers = [row[0] for row in _query(database, "SELECT po_number FROM purchase_orders")]
    assert sorted(po_numbers) == sorted(row[0] for row in products)
    stored_images = list(storage.rglob("*.png"))
    assert len(stored_images) == 1

    second = import_monthly_lists([source], database_path=database, storage_root=storage)
    assert second["products_created"] == 0
    assert second["products_existing"] == 5
    assert second["images_imported"] == 0
    assert second["images_skipped"] == 1
    assert second["tracking_created"] == 0
    assert second["product_codes_repaired"] == 0
    assert len(_query(database, "SELECT internal_id FROM products")) == 5
    assert len(_query(database, "SELECT internal_id FROM order_tracking_items")) == 5
    assert len(list(storage.rglob("*.png"))) == 1
