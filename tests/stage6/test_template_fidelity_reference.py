"""Canonical Excel reference-template fidelity gate.

The customer's exact business workbook is an external input.  Drop it at
``<HMS_QR_TEST_ROOT>/reference/excel-template.xlsx`` (or point
``HMS_QR_EXCEL_REFERENCE_TEMPLATE`` at it) and this module proves that
``TemplatePreservingExporter`` reproduces the template in every
unedited dimension: values, formulas, fonts, fills, borders, alignment,
number formats, merges, column/row geometry, defined names and image
inventory.  Without the reference file the gate skips instead of
claiming fidelity.
"""

from __future__ import annotations

import hashlib
import os
from pathlib import Path

import pytest
from openpyxl import load_workbook

from config.paths import require_test_root
from packages.excel.template_adapter import (
    TemplateCellUpdate,
    TemplateImagePlacement,
    TemplatePreservingExporter,
)
from packages.generated_assets import generated_asset_root


PNG_1X1 = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01"
    b"\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
)

PROBE_CELL = "ZZ1000"
PROBE_VALUE = "HMS-FIDELITY-PROBE"


def _reference_template() -> Path | None:
    override = os.environ.get("HMS_QR_EXCEL_REFERENCE_TEMPLATE", "").strip()
    if override:
        candidate = Path(override)
        return candidate if candidate.is_file() else None
    candidate = require_test_root() / "reference" / "excel-template.xlsx"
    return candidate if candidate.is_file() else None


def _color_signature(color) -> tuple:
    if color is None:
        return None
    return (color.rgb, color.theme, color.indexed, color.tint, color.type, color.auto)


def _side_signature(side) -> tuple:
    if side is None:
        return None
    return (side.style, _color_signature(side.color))


def _font_signature(font) -> tuple:
    return (
        font.name,
        font.charset,
        font.family,
        font.b,
        font.i,
        font.strike,
        font.outline,
        font.shadow,
        font.condense,
        _color_signature(font.color),
        font.extend,
        font.sz,
        font.u,
        font.vertAlign,
        font.scheme,
    )


def _fill_signature(fill) -> tuple:
    return (fill.patternType, _color_signature(fill.fgColor), _color_signature(fill.bgColor))


def _border_signature(border) -> tuple:
    return (
        _side_signature(border.left),
        _side_signature(border.right),
        _side_signature(border.top),
        _side_signature(border.bottom),
        _side_signature(border.diagonal),
        border.diagonalUp,
        border.diagonalDown,
        border.outline,
    )


def _alignment_signature(alignment) -> tuple:
    return (
        alignment.horizontal,
        alignment.vertical,
        alignment.wrapText,
        alignment.indent,
        alignment.shrinkToFit,
        alignment.textRotation,
        alignment.relativeIndent,
        alignment.justifyLastLine,
        alignment.readingOrder,
    )


def _protection_signature(protection) -> tuple:
    return (protection.locked, protection.hidden)


def _style_signature(cell) -> tuple:
    return (
        cell.value,
        _font_signature(cell.font),
        _fill_signature(cell.fill),
        _border_signature(cell.border),
        _alignment_signature(cell.alignment),
        cell.number_format,
        _protection_signature(cell.protection),
    )


def _sheet_signature(sheet) -> dict:
    signature: dict[str, object] = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None and cell.style == "Normal":
                continue
            signature[cell.coordinate] = _style_signature(cell)
    signature["__merged__"] = sorted(str(rng) for rng in sheet.merged_cells.ranges)
    signature["__columns__"] = {
        letter: (dimension.width, dimension.hidden, dimension.outlineLevel)
        for letter, dimension in sheet.column_dimensions.items()
    }
    signature["__rows__"] = {
        index: (dimension.height, dimension.hidden, dimension.outlineLevel)
        for index, dimension in sheet.row_dimensions.items()
    }
    signature["__images__"] = len(sheet._images)
    signature["__freeze__"] = sheet.freeze_panes
    signature["__sheet_state__"] = sheet.sheet_state
    return signature


def _workbook_signature(path: Path) -> dict:
    workbook = load_workbook(path, data_only=False)
    try:
        return {
            "sheetnames": workbook.sheetnames,
            "defined_names": sorted(workbook.defined_names),
            "sheets": {name: _sheet_signature(workbook[name]) for name in workbook.sheetnames},
        }
    finally:
        workbook.close()


def test_reference_template_export_preserves_every_unedited_dimension(tmp_path: Path):
    reference = _reference_template()
    if reference is None:
        pytest.skip("canonical Excel reference template not provided")
    assert require_test_root() in tmp_path.resolve().parents
    before_hash = hashlib.sha256(reference.read_bytes()).hexdigest()
    before = _workbook_signature(reference)

    first_sheet = before["sheetnames"][0]
    image = tmp_path / "probe.png"
    image.write_bytes(PNG_1X1)
    output = generated_asset_root() / "reference-fidelity-probe.xlsx"
    result = TemplatePreservingExporter(allowed_output_root=generated_asset_root()).export(
        template_path=reference,
        output_path=output,
        cell_updates=(TemplateCellUpdate(first_sheet, PROBE_CELL, PROBE_VALUE),),
        images=(TemplateImagePlacement(first_sheet, PROBE_CELL, image, width=48, height=48),),
    )

    assert result.source_unchanged
    assert hashlib.sha256(reference.read_bytes()).hexdigest() == before_hash
    after = _workbook_signature(output)

    assert after["sheetnames"] == before["sheetnames"]
    assert after["defined_names"] == before["defined_names"]
    for name in before["sheetnames"]:
        source_sheet = before["sheets"][name]
        exported_sheet = after["sheets"][name]
        for key, source_value in source_sheet.items():
            if key == "__images__":
                continue
            assert exported_sheet[key] == source_value, f"drift at sheet {name!r} key {key!r}"
        if name == first_sheet:
            assert exported_sheet["__images__"] == source_sheet["__images__"] + 1
            assert exported_sheet[PROBE_CELL][0] == PROBE_VALUE
            extra = set(exported_sheet) - set(source_sheet) - {PROBE_CELL}
            assert not extra, f"unexpected new cells in export: {sorted(extra)}"
        else:
            assert exported_sheet["__images__"] == source_sheet["__images__"]
